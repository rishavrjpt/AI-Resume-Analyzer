"""
Unit tests for document parsing service:
Covers TXT, PDF, and DOCX formats, empty files, corrupt data, and size limits.
"""

import os
import io
import pytest
from app.services.document_parser import DocumentParser, DocumentParseError


def test_txt_parsing_valid():
    content = "John Doe\nSummary\nExperienced backend engineer.\nSkills\nPython, Docker"
    res = DocumentParser.parse_bytes("test_resume.txt", content.encode("utf-8"))
    assert res["file_type"] == ".txt"
    assert res["line_count"] == 5
    assert "Python, Docker" in res["clean_text"]


def test_unsupported_file_extension():
    with pytest.raises(DocumentParseError) as exc_info:
        DocumentParser.parse_bytes("resume.exe", b"binary content")
    assert "Unsupported file format" in str(exc_info.value)


def test_empty_file_rejected():
    with pytest.raises(DocumentParseError) as exc_info:
        DocumentParser.parse_bytes("empty.pdf", b"")
    assert "empty" in str(exc_info.value).lower()


def test_oversized_file_rejected():
    oversized_bytes = b"0" * (6 * 1024 * 1024)  # 6 MB
    with pytest.raises(DocumentParseError) as exc_info:
        DocumentParser.parse_bytes("huge.txt", oversized_bytes)
    assert "exceeds" in str(exc_info.value).lower()


def test_corrupt_pdf_handling():
    corrupt_pdf_bytes = b"%PDF-1.4 corrupt content that cannot be parsed by pypdf"
    with pytest.raises(DocumentParseError) as exc_info:
        DocumentParser.parse_bytes("corrupt.pdf", corrupt_pdf_bytes)
    assert "Failed to parse PDF" in str(exc_info.value)


def test_docx_generation_and_parsing():
    import docx
    doc = docx.Document()
    doc.add_heading("Alice Wonder", 0)
    doc.add_paragraph("Summary: Cloud architect with AWS experience.")
    doc.add_paragraph("Skills: Terraform, Python, Kubernetes.")

    buffer = io.BytesIO()
    doc.save(buffer)
    docx_bytes = buffer.getvalue()

    res = DocumentParser.parse_bytes("alice.docx", docx_bytes)
    assert res["file_type"] == ".docx"
    assert "Alice Wonder" in res["clean_text"]
    assert "Terraform" in res["clean_text"]
